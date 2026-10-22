"""Exported reports: read them (agent 21), match them to the chart that produced them (agent 22), and compare their
figures with the chart's SQL (V5, figures match).

Qlik exports a table or chart as Excel or CSV: a header row with the dimension and measure labels, an optional
totals row ("Totals"), then one row per dimension combination. Numbers in Excel keep their value; numbers in CSV
are text in the app's number format (thousand and decimal separators, currency, %). Null dimension values show as
"-" or an empty cell.
"""
from __future__ import annotations

import csv
import datetime as dt
import io
import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from adce_core.bi_ir import Visual
from pydantic import BaseModel

from .formats import FormatError, dremio_to_strftime, qlik_to_dremio

TOTAL_LABELS = {"totals", "total", "grand total"}
NULLS = {"", "-", "null"}


@dataclass
class ExportTable:
    file: str
    sheet: str
    headers: list[str]
    rows: list[list[Any]]
    totals: list[Any] | None = None


# ---------------------------------------------------------------------------------------------- agent 21: read
def parse_number(text: Any, thousand_sep: str = ",", decimal_sep: str = ".") -> float | None:
    if isinstance(text, bool):
        return None
    if isinstance(text, int | float):
        return float(text)
    if not isinstance(text, str):
        return None
    t = text.strip()
    if not t:
        return None
    pct = t.endswith("%")
    neg = t.startswith("(") and t.endswith(")")
    t = re.sub(r"[^\d" + re.escape(thousand_sep + decimal_sep) + r"\-+eE]", "", t.strip("()%"))
    t = t.replace(thousand_sep, "") if thousand_sep else t
    t = t.replace(decimal_sep, ".") if decimal_sep != "." else t
    try:
        v = float(t)
    except ValueError:
        return None
    v = -v if neg else v
    return v / 100 if pct else v


def _rows_to_table(file: str, sheet: str, rows: list[list[Any]]) -> ExportTable | None:
    rows = [r for r in rows if any(c not in (None, "") for c in r)]
    if not rows:
        return None
    width = max(len(r) for r in rows)
    rows = [list(r) + [None] * (width - len(r)) for r in rows]
    headers = [str(c).strip() if c is not None else "" for c in rows[0]]
    body, totals = [], None
    for r in rows[1:]:
        first = str(r[0]).strip().lower() if r[0] is not None else ""
        if first in TOTAL_LABELS and totals is None:
            totals = r
        else:
            body.append(r)
    return ExportTable(file=file, sheet=sheet, headers=headers, rows=body, totals=totals)


def read_export(path: str | Path, data: bytes | None = None) -> list[ExportTable]:
    """Excel (.xlsx) or CSV / TXT. PDF and images are not read here: they go to the parse_export AI route."""
    p = Path(path)
    suffix = p.suffix.lower()
    if suffix in (".xlsx", ".xlsm"):
        from openpyxl import load_workbook

        wb = load_workbook(io.BytesIO(data) if data is not None else p, read_only=True, data_only=True)
        out = []
        for ws in wb.worksheets:
            t = _rows_to_table(p.name, ws.title, [list(r) for r in ws.iter_rows(values_only=True)])
            if t is not None:
                out.append(t)
        return out
    if suffix in (".csv", ".txt", ".tsv"):
        text = data.decode("utf-8-sig") if data is not None else p.read_text(encoding="utf-8-sig")
        try:
            dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t|")
            delim = dialect.delimiter
        except csv.Error:
            delim = ","
        t = _rows_to_table(p.name, p.stem, [list(r) for r in csv.reader(io.StringIO(text), delimiter=delim)])
        return [t] if t is not None else []
    raise ValueError(f"export format {suffix} is read by the parse_export AI route, not here")


# ---------------------------------------------------------------------------------------------- agent 21: PDF
MAX_PDF_PAGES = 20
MAX_PDF_CHARS = 60_000  # what the parse_export route is given; a longer export is cut and flagged


def pdf_pages(data: bytes) -> list[str]:
    """The text of each page of a PDF export (layout kept where the PDF has it). Scanned PDFs have no text."""
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    pages = []
    for page in reader.pages[:MAX_PDF_PAGES]:
        try:
            pages.append(page.extract_text(extraction_mode="layout") or "")
        except Exception:  # noqa: BLE001 - some PDFs fail in layout mode; plain extraction still works
            pages.append(page.extract_text() or "")
    return pages


class ParsedTable(BaseModel):
    """One table of figures the model read from a PDF export (route parse_export)."""

    title: str = ""
    page: int = 1
    headers: list[str]
    rows: list[list[str]]  # cells exactly as shown (numbers as text, in the report's format)
    totals: list[str] | None = None  # the totals row, when the report shows one


class ParsedExport(BaseModel):
    """Structured output of the route parse_export."""

    tables: list[ParsedTable] = []
    notes: str = ""


def tables_from_parsed(file: str, parsed: ParsedExport) -> list[ExportTable]:
    """The model's tables as ExportTables; a table whose rows do not fit its headers is dropped."""
    out = []
    for i, t in enumerate(parsed.tables):
        width = len(t.headers)
        if not width or any(len(r) != width for r in t.rows) or (t.totals is not None and len(t.totals) != width):
            continue
        sheet = t.title.strip() or f"page {t.page} table {i + 1}"
        out.append(ExportTable(file=file, sheet=sheet, headers=[h.strip() for h in t.headers],
                               rows=[list(r) for r in t.rows], totals=list(t.totals) if t.totals else None))
    return out


# ---------------------------------------------------------------------------------------------- agent 22: map
def _norm(s: str) -> str:
    return re.sub(r"\s+", " ", s.strip().lower())


@dataclass
class ExportMatch:
    export: ExportTable
    visual_id: str | None
    columns: dict[int, tuple[str, int]] = field(default_factory=dict)  # export column -> (dimension|measure, index)
    confidence: str = "none"  # high | medium | none
    reason: str = ""


def map_export(e: ExportTable, visuals: list[Visual]) -> ExportMatch:
    """Deterministic matching on labels: every export column must match a dimension or measure label of the visual,
    and every dimension of the visual must be present. The file name or sheet naming the visual's title breaks
    ties. Anything ambiguous is returned without a visual, for the map_metric AI route or a reviewer."""
    heads = [_norm(h) for h in e.headers]
    candidates: list[tuple[int, ExportMatch]] = []
    for v in visuals:
        dims = [_norm(d.label) for d in v.dimensions]
        meas = [_norm(m.label) for m in v.measures]
        cols: dict[int, tuple[str, int]] = {}
        for i, h in enumerate(heads):
            if h in dims:
                cols[i] = ("dimension", dims.index(h))
            elif h in meas:
                cols[i] = ("measure", meas.index(h))
        if len(cols) != len(heads) or {i for k, i in cols.values() if k == "dimension"} != set(range(len(dims))):
            continue
        named = any(_norm(v.title) and _norm(v.title) in _norm(x) for x in (Path(e.file).stem, e.sheet))
        candidates.append((2 if named else 1, ExportMatch(e, v.id, cols, "high" if named else "medium",
                                                         "labels and title" if named else "labels")))
    if not candidates:
        return ExportMatch(e, None, reason="no visual has exactly these columns")
    candidates.sort(key=lambda c: -c[0])
    if len(candidates) > 1 and candidates[0][0] == candidates[1][0]:
        return ExportMatch(e, None, reason="several visuals have these columns")
    return candidates[0][1]


# ---------------------------------------------------------------------------------------------- V5: compare
@dataclass
class FigureDiff:
    row: tuple[Any, ...]
    measure: str
    expected: Any
    actual: Any


@dataclass
class FigureCheck:
    visual_id: str
    export: str
    rows_expected: int
    rows_actual: int
    figures: int = 0
    matched: int = 0
    diffs: list[FigureDiff] = field(default_factory=list)
    missing_rows: list[tuple[Any, ...]] = field(default_factory=list)  # in the export, not produced by the SQL
    extra_rows: list[tuple[Any, ...]] = field(default_factory=list)  # produced by the SQL, not in the export
    totals_checked: int = 0
    totals_matched: int = 0

    @property
    def ok(self) -> bool:
        return (self.figures > 0 and self.matched == self.figures and not self.missing_rows and not self.extra_rows
                and self.totals_matched == self.totals_checked)


def _key(value: Any, date_fmt: str | None) -> Any:
    if value is None:
        return None
    if isinstance(value, dt.datetime | dt.date):
        if date_fmt:
            try:
                return value.strftime(dremio_to_strftime(qlik_to_dremio(date_fmt)))
            except FormatError:
                pass
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    s = str(value).strip()
    if s.lower() in NULLS:
        return None
    n = parse_number(s) if re.fullmatch(r"-?[\d.,]+", s) else None
    return str(int(n)) if n is not None and n.is_integer() else s


def _close(a: Any, e: Any, thousand: str, decimal: str) -> bool:
    if a is None and (e is None or (isinstance(e, str) and e.strip().lower() in NULLS)):
        return True
    ea, ee = parse_number(a, thousand, decimal), parse_number(e, thousand, decimal)
    if ea is None or ee is None:
        return str(a).strip() == str(e).strip()
    return abs(ea - ee) <= max(0.005, 1e-9 * abs(ee))  # exports often round to 2 decimals


def compare(match: ExportMatch, actual_columns: list[str], actual_rows: list[tuple[Any, ...]],
            actual_totals: tuple[Any, ...] | None = None, *, thousand_sep: str = ",", decimal_sep: str = ".",
            date_format: str | None = None) -> FigureCheck:
    """actual_*: the chart SQL's result (dimension columns then measure columns, in the visual's order)."""
    e = match.export
    dim_cols = sorted((i for i, (k, _) in match.columns.items() if k == "dimension"),
                      key=lambda i: match.columns[i][1])
    meas_cols = [(i, match.columns[i][1]) for i, (k, _) in match.columns.items() if k == "measure"]
    n_dims = len(dim_cols)
    check = FigureCheck(visual_id=match.visual_id or "", export=e.file, rows_expected=len(e.rows),
                        rows_actual=len(actual_rows))
    actual = {tuple(_key(v, date_format) for v in r[:n_dims]): r for r in actual_rows}
    seen = set()
    for row in e.rows:
        k = tuple(_key(row[i], date_format) for i in dim_cols)
        a = actual.get(k)
        if a is None:
            check.missing_rows.append(k)
            continue
        seen.add(k)
        for i, mi in meas_cols:
            check.figures += 1
            if _close(a[n_dims + mi], row[i], thousand_sep, decimal_sep):
                check.matched += 1
            elif len(check.diffs) < 50:
                check.diffs.append(FigureDiff(k, actual_columns[n_dims + mi], row[i], a[n_dims + mi]))
    check.extra_rows = [k for k in actual if k not in seen]
    if e.totals is not None and actual_totals is not None:
        for i, mi in meas_cols:
            if e.totals[i] in (None, ""):
                continue
            check.totals_checked += 1
            if _close(actual_totals[mi], e.totals[i], thousand_sep, decimal_sep):
                check.totals_matched += 1
            elif len(check.diffs) < 50:
                check.diffs.append(FigureDiff(("Totals",), actual_columns[n_dims + mi], e.totals[i],
                                              actual_totals[mi]))
    return check
