"""The migration pack (ADR-Q7): the files agent 20 publishes for one application, built in memory.

```
<app>__<reload_time>__pack-v<N>/
  manifest.json                 app, inputs, tool, model and prompt versions, run id, content fingerprint
  sql/tables/<qlik_table>.sql   Dremio SQL per Qlik table
  sql/measures/<visual>.sql     SQL per chart (its measures over its dimensions)
  lineage/lineage.json          OpenLineage-shaped jobs; lineage.md the readable summary with the QVD chain
  validation/report.md          ladder level per table and figure, failed comparisons, drift; report.xlsx the same
  review/untranslated.csv       statements not translated, with spans
  review/decisions.csv          reviewer (and auto-reviewer) decisions applied
  review/open-items.md          what is not resolved: untranslated, inputs, unmapped sources, questions, RLS
```

The pack is reproducible: everything here is a pure function of its inputs (no clock, sorted orders, a fixed date
inside the xlsx), so the same inputs and versions produce the same bytes. `fingerprint` hashes every file but the
manifest; the Publisher reuses an existing version with the same fingerprint instead of writing a new one.
"""
from __future__ import annotations

import csv
import hashlib
import io
import json
import zipfile
from datetime import datetime
from typing import Any

from adce_reversesql.reconcile import DRIFT

FIXED_TIME = datetime(2000, 1, 1)  # inside report.xlsx, so the bytes do not depend on when it was written


def fingerprint(files: dict[str, bytes]) -> str:
    h = hashlib.sha256()
    for path in sorted(files):
        if path != "manifest.json":
            h.update(path.encode() + b"\0" + hashlib.sha256(files[path]).digest())
    return h.hexdigest()


def pack_name(app_id: str, reload_time: str | None, version: int) -> str:
    return f"{app_id}__{reload_time or 'no-reload-log'}__pack-v{version}"


def _csv(header: list[str], rows: list[list[Any]]) -> bytes:
    buf = io.StringIO()
    w = csv.writer(buf, lineterminator="\n")
    w.writerow(header)
    w.writerows(rows)
    return buf.getvalue().encode()


# ---------------------------------------------------------------- validation
def validation_md(report: list[dict[str, Any]], figures: list[dict[str, Any]] | None = None) -> str:
    lines = ["# Validation report", "",
             "Levels: V0 parsed · V1 checked · V2 executes · V3 row counts · V4 data · V5 figures match exports.",
             "Venue `local` means DuckDB over QVD and source extracts; results are re-earned in Dremio in phase PD.",
             "", "| Table | Level | Venue | Comparisons | Failed |", "|---|---|---|---|---|"]
    for r in report:
        comps = r["comparisons"]
        bad = [c for c in comps if not c["ok"]]
        level = r["level"] if r.get("exact", True) else f"{r['level']} (within tolerance, drift)"
        lines.append(f"| {r['table']} | {level} | {r['venue']} | {len(comps)} | {len(bad)} |")
    for r in report:
        bad = [c for c in r["comparisons"] if not c["ok"]]
        failed_rules = [x for x in r["rules"] if not x["ok"]]
        drift = [c for c in r["comparisons"] if c["ok"] and c.get("basis") == DRIFT]
        if bad or failed_rules or r["errors"] or drift:
            lines += ["", f"## {r['table']}"]
            lines += [f"- rule `{x['name']}`: {x['message']}" for x in failed_rules]
            lines += [f"- {c['check']} {c['subject']}: expected {c['expected']}, got {c['actual']}" for c in bad]
            lines += [f"- within tolerance (drift): {c['check']} {c['subject']}: expected {c['expected']}, "
                      f"got {c['actual']}" for c in drift]
            lines += [f"- error: {e}" for e in r["errors"]]
    if figures:
        lines += ["", "## Charts against exported reports", "",
                  "| Export | Visual | Level | Figures | Totals | Rows (SQL / export) |", "|---|---|---|---|---|---|"]
        for f in figures:
            lines.append(f"| {f['export']} [{f['sheet']}] | {f['visual'] or '-'} | {f['level']} | "
                         f"{f.get('matched', 0)}/{f.get('figures', 0)} | {f.get('totals_matched', 0)}/"
                         f"{f.get('totals', 0)} | {f.get('rows_actual', '-')} / {f.get('rows_expected', '-')} |")
        for f in figures:
            if f["level"] != "V5":
                lines += ["", f"### {f['export']} [{f['sheet']}]"]
                lines += [f"- {p}" for p in f.get("problems", [])]
                lines += [f"- {d['row']} {d['measure']}: expected {d['expected']}, got {d['actual']}"
                          for d in f.get("diffs", [])]
    return "\n".join(lines) + "\n"


def _cell(v: Any) -> Any:
    return v if isinstance(v, (int, float, str)) or v is None else json.dumps(v, default=str, sort_keys=True)


def validation_xlsx(report: list[dict[str, Any]], figures: list[dict[str, Any]] | None = None) -> bytes:
    """The validation report as a workbook: Tables, Comparisons, Rules, Figures. Byte-identical for the same
    report (fixed document dates and zip entry times)."""
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "Tables"
    ws.append(["Table", "Level", "Venue", "Exact", "Comparisons", "Failed", "Errors"])
    for r in report:
        comps = r["comparisons"]
        ws.append([r["table"], r["level"], r["venue"], "yes" if r.get("exact", True) else "within tolerance (drift)",
                   len(comps), sum(1 for c in comps if not c["ok"]), "; ".join(r["errors"])])
    wc = wb.create_sheet("Comparisons")
    wc.append(["Table", "Check", "Subject", "Expected", "Actual", "OK", "Basis"])
    for r in report:
        for c in r["comparisons"]:
            wc.append([r["table"], c["check"], c["subject"], _cell(c["expected"]), _cell(c["actual"]),
                       "yes" if c["ok"] else "no", c.get("basis", "")])
    wr = wb.create_sheet("Rules")
    wr.append(["Table", "Rule", "OK", "Message"])
    for r in report:
        for x in r["rules"]:
            wr.append([r["table"], x["name"], "yes" if x["ok"] else "no", x.get("message", "")])
    wf = wb.create_sheet("Figures")
    wf.append(["Export", "Sheet", "Visual", "Level", "Figures matched", "Figures", "Totals matched", "Totals"])
    for f in figures or []:
        wf.append([f["export"], f["sheet"], f.get("visual") or "", f["level"], f.get("matched", 0),
                   f.get("figures", 0), f.get("totals_matched", 0), f.get("totals", 0)])
    wb.properties.created = FIXED_TIME
    wb.properties.modified = FIXED_TIME
    wb.properties.creator = "adce"
    raw = io.BytesIO()
    wb.save(raw)
    return _stable_zip(raw.getvalue())


def _stable_zip(data: bytes) -> bytes:
    """Rewrite a zip with fixed entry times (openpyxl stamps the current time, and docProps/core.xml the save time)."""
    out = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(data)) as src, zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as dst:
        for info in sorted(src.infolist(), key=lambda i: i.filename):
            body = src.read(info.filename)
            if info.filename == "docProps/core.xml":
                body = _fixed_core(body)
            z = zipfile.ZipInfo(info.filename, date_time=(2000, 1, 1, 0, 0, 0))
            z.compress_type = zipfile.ZIP_DEFLATED
            dst.writestr(z, body)
    return out.getvalue()


def _fixed_core(xml: bytes) -> bytes:
    import re

    return re.sub(rb"(<dcterms:(?:created|modified)[^>]*>)[^<]*(<)", rb"\g<1>2000-01-01T00:00:00Z\g<2>", xml)


# ---------------------------------------------------------------- lineage
def lineage_md(app_id: str, lin: dict[str, Any]) -> str:
    """Readable lineage: what each table is built from and where it is stored, then the QVD chain."""
    lines = [f"# Lineage of {app_id}", "", "Dremio datasets → Qlik tables → QVDs. Machine-readable: lineage.json "
             "(OpenLineage-shaped).", "", "| Table | Built from | Stored to |", "|---|---|---|"]
    qvd_made: dict[str, str] = {}
    qvd_read: dict[str, list[str]] = {}
    for job in sorted(lin["jobs"], key=lambda j: j["job"]["name"]):
        table = job["job"]["name"].removeprefix("load:")
        ins = [f"{i['name']} ({i['namespace'].split(':')[0]})" for i in job["inputs"]]
        outs = [o["name"] for o in job["outputs"] if o["namespace"] == "qvd://"]
        for o in outs:
            qvd_made[o] = table
        for i in job["inputs"]:
            if i["namespace"] == "qvd://":
                qvd_read.setdefault(i["name"], []).append(table)
        dropped = " (dropped)" if job.get("dropped") else ""
        lines.append(f"| {table}{dropped} | {', '.join(ins) or '-'} | {', '.join(outs) or '-'} |")
    if qvd_made or qvd_read:
        lines += ["", "## QVD chain", ""]
        for q in sorted(set(qvd_made) | set(qvd_read)):
            producer = qvd_made.get(q)
            by = f"made by {producer}" if producer else "made outside this application (producer not supplied)"
            readers = ", ".join(sorted(qvd_read.get(q, []))) or "not read here"
            lines.append(f"- `{q}`: {by}; read by {readers}")
    return "\n".join(lines) + "\n"


# ---------------------------------------------------------------- review
def untranslated_csv(loaded: Any, charts: list[Any]) -> bytes:
    rows: list[list[Any]] = []
    for name in sorted(loaded.translation.tables):
        for u in loaded.translation.tables[name].untranslated:
            rows.append([name, u.statement, u.span, u.tier, u.reason, u.text])
    for c in sorted(charts, key=lambda c: c.visual_id):
        if c.untranslated:
            rows.append([f"chart:{c.visual_id}", "", "", c.tier, "; ".join(c.untranslated), c.title or ""])
    return _csv(["table", "statement", "span", "tier", "reason", "text"], rows)


def decisions_csv(decisions: list[Any]) -> bytes:
    """Decisions applied: facts decided by a reviewer (`_decided_by`) and the auto-reviewer's acceptances."""
    rows = []
    for f in sorted(decisions, key=lambda f: (f.kind, f.subject, f.id)):
        v = dict(f.value)
        by = v.pop("_decided_by", None) or f.agent or ""
        decision = v.pop("_decision", None) or ("accepted" if f.kind == "auto_accepted" else f.status.value)
        shown = {k: x for k, x in v.items() if not k.startswith("_")}
        rows.append([f.kind, f.subject, decision, by, json.dumps(shown, sort_keys=True, default=str)[:2000], f.id])
    return _csv(["kind", "subject", "decision", "by", "value", "fact_id"], rows)


def open_items_md(loaded: Any, charts: list[Any], *, unmapped: list[str], missing_producers: list[str],
                  not_approved: list[str]) -> str:
    lines = ["# Open items", "", "What is not resolved in this pack. Each item names what would resolve it.", ""]

    def section(title: str, items: list[str]) -> None:
        nonlocal lines
        lines += [f"## {title}", ""] + ([f"- {i}" for i in items] or ["- none"]) + [""]

    section("Statements not translated (provide SQL in the review queue, or accept as a known gap)",
            [f"{n}: statement {u.statement} ({u.span}) - {u.reason} [{u.tier}]"
             for n in sorted(loaded.translation.tables) for u in loaded.translation.tables[n].untranslated])
    section("Charts not translated", [f"{c.visual_id} {c.title!r}: {'; '.join(c.untranslated)} [{c.tier}]"
                                      for c in sorted(charts, key=lambda c: c.visual_id) if c.untranslated])
    section("Sources without a Dremio dataset (map them on the Sources page)", sorted(unmapped))
    section("QVDs read whose producer application was not supplied (supply it, or map the QVD to a dataset)",
            sorted(missing_producers))
    section("Owner questions (variables the script could not resolve)",
            sorted(i.message for i in loaded.app.issues if i.code == "unresolved_variable"))
    section("Tables whose SQL is not approved yet", sorted(not_approved))
    if loaded.app.section_access:
        section("Row-level security (Section Access): requirements to implement in Dremio",
                [f"{a.name}: data reduced on {', '.join(a.reduction_fields) or 'no field'} (access fields: "
                 f"{', '.join(a.fields)})" for a in loaded.app.section_access])
    section("Other issues found while reading the script",
            sorted(f"{i.code}: {i.message}" for i in loaded.app.issues if i.code != "unresolved_variable"))
    return "\n".join(lines)


def versions() -> dict[str, Any]:
    """Tool versions (ADCE packages installed) and the prompt version of every model route."""
    from importlib.metadata import PackageNotFoundError, version

    from adce_gateway.routes import ROUTES

    pkgs = {}
    for p in ("adce-core", "adce-qlik", "adce-reversesql", "adce-gateway", "adce-agents-reversesql",
              "adce-agents-bi-adapters", "sqlglot", "duckdb"):
        try:
            pkgs[p] = version(p)
        except PackageNotFoundError:
            continue
    return {"software": pkgs, "routes": {n: r.prompt_version for n, r in sorted(ROUTES.items())}}
